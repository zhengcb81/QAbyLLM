#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档处理器
支持多种文档格式的扩展处理
"""

import os
import re
from typing import Dict, List, Any, Optional, Tuple
from pathlib import Path
import logging
from enum import Enum
import tempfile

logger = logging.getLogger(__name__)

class DocumentFormat(Enum):
    """文档格式枚举"""
    TXT = "txt"
    MD = "md"
    JSON = "json"
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    PPTX = "pptx"
    HTML = "html"
    EPUB = "epub"
    IMAGE = "image"  # 图片格式（需要OCR）

class DocumentProcessor:
    """文档处理器"""
    
    def __init__(self, config: Optional[Dict[str, Any]] = None):
        self.config = config or {}
        self.supported_formats = self._get_supported_formats()
        self.ocr_enabled = self.config.get('ocr', {}).get('enabled', False)
        self.table_processing_enabled = self.config.get('tables', {}).get('enabled', True)
    
    def _get_supported_formats(self) -> List[DocumentFormat]:
        """获取支持的格式列表"""
        formats = [
            DocumentFormat.TXT,
            DocumentFormat.MD,
            DocumentFormat.JSON,
            DocumentFormat.PDF,
            DocumentFormat.DOCX,
            DocumentFormat.XLSX,
            DocumentFormat.PPTX,
            DocumentFormat.HTML
        ]
        
        # 检查可选依赖
        try:
            import ebooklib
            formats.append(DocumentFormat.EPUB)
        except ImportError:
            logger.warning("EPUB支持不可用，请安装ebooklib")
        
        if self.ocr_enabled:
            formats.append(DocumentFormat.IMAGE)
        
        return formats
    
    def get_file_format(self, file_path: str) -> Optional[DocumentFormat]:
        """获取文件格式"""
        try:
            ext = Path(file_path).suffix.lower().lstrip('.')
            
            # 图片格式检测
            if ext in ['jpg', 'jpeg', 'png', 'bmp', 'tiff', 'tif']:
                return DocumentFormat.IMAGE
            
            # 其他格式
            for format_enum in DocumentFormat:
                if format_enum.value == ext:
                    return format_enum
            
            return None
            
        except Exception as e:
            logger.error(f"获取文件格式失败 {file_path}: {e}")
            return None
    
    def is_supported_format(self, file_path: str) -> bool:
        """检查文件格式是否支持"""
        file_format = self.get_file_format(file_path)
        return file_format in self.supported_formats if file_format else False
    
    def read_document(self, file_path: str) -> Dict[str, Any]:
        """读取文档内容"""
        file_format = self.get_file_format(file_path)
        
        if not file_format:
            return self._create_error_result(f"不支持的文件格式: {file_path}")
        
        if file_format not in self.supported_formats:
            return self._create_error_result(f"格式不支持: {file_format.value}")
        
        try:
            content = ""
            metadata = {
                'file_path': file_path,
                'file_name': Path(file_path).name,
                'file_format': file_format.value,
                'file_size': os.path.getsize(file_path),
                'processing_time': 0
            }
            
            import time
            start_time = time.time()
            
            if file_format == DocumentFormat.TXT:
                content = self._read_text(file_path)
            elif file_format == DocumentFormat.MD:
                content = self._read_markdown(file_path)
            elif file_format == DocumentFormat.JSON:
                content = self._read_json(file_path)
            elif file_format == DocumentFormat.PDF:
                content = self._read_pdf(file_path)
            elif file_format == DocumentFormat.DOCX:
                content = self._read_docx(file_path)
            elif file_format == DocumentFormat.XLSX:
                content = self._read_excel(file_path)
            elif file_format == DocumentFormat.PPTX:
                content = self._read_pptx(file_path)
            elif file_format == DocumentFormat.HTML:
                content = self._read_html(file_path)
            elif file_format == DocumentFormat.EPUB:
                content = self._read_epub(file_path)
            elif file_format == DocumentFormat.IMAGE:
                content = self._read_image_with_ocr(file_path)
            
            metadata['processing_time'] = time.time() - start_time
            metadata['content_length'] = len(content)
            
            return {
                'success': True,
                'content': content,
                'metadata': metadata,
                'error': None
            }
            
        except Exception as e:
            logger.error(f"读取文档失败 {file_path}: {e}")
            return self._create_error_result(f"读取失败: {str(e)}")
    
    def _create_error_result(self, error_msg: str) -> Dict[str, Any]:
        """创建错误结果"""
        return {
            'success': False,
            'content': '',
            'metadata': {},
            'error': error_msg
        }
    
    def _read_text(self, file_path: str) -> str:
        """读取文本文件"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def _read_markdown(self, file_path: str) -> str:
        """读取Markdown文件"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def _read_json(self, file_path: str) -> str:
        """读取JSON文件"""
        import json
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return json.dumps(data, ensure_ascii=False, indent=2)
    
    def _read_pdf(self, file_path: str) -> str:
        """读取PDF文件"""
        try:
            # 尝试使用PyPDF2
            import PyPDF2
            
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
                
        except ImportError:
            # 回退到pdfminer
            try:
                from pdfminer.high_level import extract_text
                return extract_text(file_path)
            except ImportError:
                raise Exception("需要安装PyPDF2或pdfminer.six来读取PDF文件")
    
    def _read_docx(self, file_path: str) -> str:
        """读取Word文档"""
        from docx import Document
        
        doc = Document(file_path)
        text = ""
        
        # 读取段落
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # 读取表格
        if self.table_processing_enabled:
            for table in doc.tables:
                text += self._extract_table_text(table) + "\n"
        
        return text
    
    def _read_excel(self, file_path: str) -> str:
        """读取Excel文件"""
        import openpyxl
        
        workbook = openpyxl.load_workbook(file_path)
        text = ""
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text += f"工作表: {sheet_name}\n"
            
            # 读取表格数据
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                text += row_text + "\n"
            
            text += "\n"
        
        return text
    
    def _read_pptx(self, file_path: str) -> str:
        """读取PowerPoint文件"""
        from pptx import Presentation
        
        prs = Presentation(file_path)
        text = ""
        
        for slide_number, slide in enumerate(prs.slides, 1):
            text += f"幻灯片 {slide_number}:\n"
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text += shape.text + "\n"
            
            text += "\n"
        
        return text
    
    def _read_html(self, file_path: str) -> str:
        """读取HTML文件"""
        from bs4 import BeautifulSoup
        
        with open(file_path, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f.read(), 'html.parser')
            
            # 移除脚本和样式
            for script in soup(["script", "style"]):
                script.decompose()
            
            # 获取文本内容
            text = soup.get_text()
            
            # 清理空白字符
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return text
    
    def _read_epub(self, file_path: str) -> str:
        """读取EPUB文件"""
        from ebooklib import epub
        from bs4 import BeautifulSoup
        
        book = epub.read_epub(file_path)
        text = ""
        
        for item in book.get_items():
            if item.get_type() == epub.EpubHtml:
                soup = BeautifulSoup(item.get_content(), 'html.parser')
                
                # 移除脚本和样式
                for script in soup(["script", "style"]):
                    script.decompose()
                
                # 获取文本内容
                chapter_text = soup.get_text()
                
                # 清理空白字符
                lines = (line.strip() for line in chapter_text.splitlines())
                chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
                chapter_text = '\n'.join(chunk for chunk in chunks if chunk)
                
                text += chapter_text + "\n\n"
        
        return text
    
    def _read_image_with_ocr(self, file_path: str) -> str:
        """使用OCR读取图片"""
        if not self.ocr_enabled:
            return "OCR功能未启用"
        
        try:
            import pytesseract
            from PIL import Image
            
            # 打开图片
            image = Image.open(file_path)
            
            # 使用OCR提取文本
            text = pytesseract.image_to_string(image, lang='chi_sim+eng')
            
            return text
            
        except ImportError:
            raise Exception("需要安装pytesseract和Pillow来使用OCR功能")
        except Exception as e:
            raise Exception(f"OCR处理失败: {e}")
    
    def _extract_table_text(self, table) -> str:
        """提取表格文本"""
        text = "表格:\n"
        
        for row in table.rows:
            row_text = "| "
            for cell in row.cells:
                row_text += cell.text.replace("\n", " ").strip() + " | "
            text += row_text + "\n"
        
        return text + "\n"
    
    def batch_process(self, folder_path: str) -> List[Dict[str, Any]]:
        """批量处理文件夹中的文档"""
        results = []
        
        if not os.path.exists(folder_path):
            logger.error(f"文件夹不存在: {folder_path}")
            return results
        
        supported_extensions = [f.value for f in self.supported_formats]
        
        for root, dirs, files in os.walk(folder_path):
            for file in files:
                file_ext = Path(file).suffix.lower().lstrip('.')
                
                # 检查图片格式
                if file_ext in ['jpg', 'jpeg', 'png', 'bmp', 'tiff', 'tif']:
                    file_ext = 'image'
                
                if file_ext in supported_extensions:
                    file_path = os.path.join(root, file)
                    result = self.read_document(file_path)
                    results.append(result)
        
        return results
    
    def get_processing_stats(self, results: List[Dict[str, Any]]) -> Dict[str, Any]:
        """获取处理统计信息"""
        total_files = len(results)
        successful_files = sum(1 for r in results if r['success'])
        failed_files = total_files - successful_files
        
        total_size = sum(r['metadata'].get('file_size', 0) for r in results if r['success'])
        total_content = sum(r['metadata'].get('content_length', 0) for r in results if r['success'])
        total_time = sum(r['metadata'].get('processing_time', 0) for r in results if r['success'])
        
        return {
            'total_files': total_files,
            'successful_files': successful_files,
            'failed_files': failed_files,
            'success_rate': successful_files / total_files if total_files > 0 else 0,
            'total_size_bytes': total_size,
            'total_content_length': total_content,
            'total_processing_time_seconds': total_time,
            'avg_processing_time_seconds': total_time / successful_files if successful_files > 0 else 0
        }

# 测试代码
if __name__ == "__main__":
    # 创建文档处理器
    processor = DocumentProcessor({
        'ocr': {'enabled': False},
        'tables': {'enabled': True}
    })
    
    # 测试支持的格式
    print("=== 支持的格式 ===")
    for fmt in processor.supported_formats:
        print(f"- {fmt.value}")
    
    # 测试文件格式检测
    test_files = [
        "document.txt",
        "report.pdf", 
        "data.xlsx",
        "presentation.pptx",
        "image.jpg"
    ]
    
    print("\n=== 文件格式检测 ===")
    for file in test_files:
        fmt = processor.get_file_format(file)
        supported = processor.is_supported_format(file)
        print(f"{file}: {fmt.value if fmt else '未知'} ({'支持' if supported else '不支持'})")
    
    # 测试文档读取（需要实际文件）
    print("\n=== 文档读取测试 ===")
    # 这里需要实际文件路径进行测试
    # result = processor.read_document("test.txt")
    # print(f"成功: {result['success']}, 长度: {len(result['content'])}")